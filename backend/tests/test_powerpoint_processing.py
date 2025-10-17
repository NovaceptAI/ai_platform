#!/usr/bin/env python3
"""
Test script for PowerPoint processing
"""

import sys
import os
sys.path.append('/home/azureuser/ai_platform/backend')

from app.utils.file_utils import extract_text_from_presentation, detect_file_type

def test_powerpoint_processing(pptx_path):
    """Test PowerPoint processing"""
    print(f"Testing PowerPoint processing for: {pptx_path}")
    
    # Check if file exists
    if not os.path.exists(pptx_path):
        print(f"❌ File not found: {pptx_path}")
        return
    
    # Detect file type
    try:
        file_type = detect_file_type(pptx_path)
        print(f"✓ Detected file type: {file_type}")
    except Exception as e:
        print(f"❌ File type detection failed: {e}")
        return
    
    # Extract text from presentation
    try:
        # Make a copy since the function deletes the file
        import shutil
        temp_path = f"/tmp/test_pptx_{os.path.basename(pptx_path)}"
        shutil.copy2(pptx_path, temp_path)
        
        text_chunks = extract_text_from_presentation(temp_path)
        print(f"✓ Text extraction successful!")
        print(f"✓ Number of chunks: {len(text_chunks)}")
        
        for i, chunk in enumerate(text_chunks, 1):
            print(f"\n--- Chunk {i} ---")
            print(chunk[:300] + "..." if len(chunk) > 300 else chunk)
            
    except Exception as e:
        print(f"❌ Text extraction failed: {e}")

def create_sample_pptx():
    """Create a sample PowerPoint file for testing"""
    try:
        from pptx import Presentation
        from pptx.util import Inches
        
        prs = Presentation()
        
        # Slide 1
        slide1 = prs.slides.add_slide(prs.slide_layouts[1])
        slide1.shapes.title.text = "Sample Presentation"
        slide1.shapes.placeholders[1].text = "This is a test presentation created for file processing validation."
        
        # Slide 2
        slide2 = prs.slides.add_slide(prs.slide_layouts[1])
        slide2.shapes.title.text = "Technical Features"
        slide2.shapes.placeholders[1].text = """• PowerPoint text extraction
• Multi-slide processing
• Token-based chunking
• Automatic cleanup"""
        
        # Slide 3
        slide3 = prs.slides.add_slide(prs.slide_layouts[1])
        slide3.shapes.title.text = "Implementation Details"
        slide3.shapes.placeholders[1].text = """The system processes PowerPoint files by:
1. Reading slide content
2. Extracting text from shapes
3. Chunking based on token limits
4. Returning structured text data"""
        
        sample_path = "/tmp/sample_presentation.pptx"
        prs.save(sample_path)
        print(f"✓ Created sample PowerPoint: {sample_path}")
        return sample_path
        
    except Exception as e:
        print(f"❌ Failed to create sample PowerPoint: {e}")
        return None

if __name__ == "__main__":
    print("=== PowerPoint Processing Test ===")
    
    # Create a sample file
    sample_file = create_sample_pptx()
    
    # Test with sample file
    if sample_file:
        test_powerpoint_processing(sample_file)
    
    # You can also test with your own files
    test_files = [
        "/path/to/your/test/presentation.pptx"
    ]
    
    for test_file in test_files:
        if os.path.exists(test_file):
            test_powerpoint_processing(test_file)
        else:
            print(f"Skipping {test_file} (file not found)")
    
    print("\n=== Test Complete ===")
